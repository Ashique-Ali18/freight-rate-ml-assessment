from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path


# ============================================================
# Configuration
# ============================================================

ROOT = Path(__file__).resolve().parent.parent
OUTPUT = ROOT / "reports" / "Spotter_ML_Assessment_Report.docx"
CHART = ROOT / "scorer_results" / "candidate_december.png"


# ============================================================
# Helpers
# ============================================================

def set_cell_shading(cell, fill):
    """Set table-cell background color."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(9)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_text(cell, header, bold=True)
        set_cell_shading(cell, "D9EAF7")

    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)

    doc.add_paragraph()
    return table


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_after = Pt(6)

    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)

    return p


def add_metric_box(doc, label, value):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    set_cell_text(table.cell(0, 0), label, bold=True)
    set_cell_text(table.cell(0, 1), value, bold=True)

    set_cell_shading(table.cell(0, 0), "EAF2F8")
    set_cell_shading(table.cell(0, 1), "F7F9FA")

    doc.add_paragraph()


# ============================================================
# Create document
# ============================================================

doc = Document()

section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)


# ============================================================
# Global font
# ============================================================

styles = doc.styles

styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10)

styles["Title"].font.name = "Aptos Display"
styles["Title"].font.size = Pt(26)
styles["Title"].font.bold = True

styles["Heading 1"].font.name = "Aptos Display"
styles["Heading 1"].font.size = Pt(17)
styles["Heading 1"].font.bold = True

styles["Heading 2"].font.name = "Aptos Display"
styles["Heading 2"].font.size = Pt(13)
styles["Heading 2"].font.bold = True


# ============================================================
# Cover Page
# ============================================================

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.space_after = Pt(20)

run = p.add_run("MACHINE LEARNING ENGINEER ASSESSMENT")
run.bold = True
run.font.size = Pt(12)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

run = p.add_run("Freight Rate Prediction")
run.bold = True
run.font.size = Pt(28)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

run = p.add_run("Technical Assessment Report")
run.font.size = Pt(18)

doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

run = p.add_run("Spotter")
run.bold = True
run.font.size = Pt(16)

doc.add_paragraph()
doc.add_paragraph()

add_metric_box(doc, "Development Dataset", "48,000 labeled loads")
add_metric_box(doc, "Final Validation Dataset", "12,000 loads")
add_metric_box(doc, "Primary Validation MAE", "$103.75")
add_metric_box(doc, "Primary Validation RMSE", "$648.45")
add_metric_box(doc, "Primary Validation R²", "0.8200")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(20)

run = p.add_run("Prepared as part of the Machine Learning Engineer assessment")
run.italic = True
run.font.size = Pt(10)

doc.add_page_break()


# ============================================================
# 1. Executive Summary
# ============================================================

add_heading(doc, "1. Executive Summary")

doc.add_paragraph(
    "This project develops a machine learning regression pipeline for predicting "
    "freight load rates from historical shipment information. The development "
    "dataset contains 48,000 labeled loads covering January through October 2025, "
    "while the final validation dataset contains 12,000 future loads from "
    "November and December 2025."
)

doc.add_paragraph(
    "The analysis began with exploratory data analysis and data-quality checks, "
    "followed by feature evaluation, model comparison, hyperparameter tuning, "
    "and time-aware validation. Because the final validation period occurs after "
    "the labeled development period, chronological validation was preferred over "
    "a random train/test split."
)

doc.add_paragraph(
    "GradientBoostingRegressor with Huber loss produced the strongest overall "
    "validation performance among the evaluated models. The final model achieved "
    "a mean absolute error of $103.75, RMSE of $648.45, and R² of 0.8200 on the "
    "October 2025 holdout."
)

doc.add_paragraph(
    "The provided scoring script successfully validated all 12,000 final "
    "predictions and all 31 required December predictions."
)


# ============================================================
# 2. Problem Definition
# ============================================================

add_heading(doc, "2. Problem Definition")

doc.add_paragraph(
    "The objective is to estimate the posted freight rate for an unseen load "
    "using historical shipment, route, equipment, geographic, and operational "
    "information."
)

add_heading(doc, "Target", level=2)

add_code(doc, "posted_rate")

add_heading(doc, "Final Prediction Format", level=2)

add_code(doc, "load_id,predicted_rate")

doc.add_paragraph(
    "The final validation output contains exactly 12,000 rows and follows the "
    "required load ID and prediction format."
)


# ============================================================
# 3. Dataset Overview
# ============================================================

add_heading(doc, "3. Dataset Overview")

add_table(
    doc,
    ["Dataset", "Rows", "Time Period", "Purpose"],
    [
        ["Development", "48,000", "Jan–Oct 2025", "Model training and validation"],
        ["Validation", "12,000", "Nov–Dec 2025", "Final prediction"],
        ["December inputs", "31", "Dec 2025", "Required fixed scenario"],
    ],
)

doc.add_paragraph(
    "The development data contains route, geographic, equipment, shipment, "
    "date, market, quote, and target information."
)

add_heading(doc, "Main Features", level=2)

for feature in [
    "pickup and delivery locations",
    "pickup and delivery latitude/longitude",
    "distance",
    "equipment",
    "weight",
    "date",
    "market_index",
    "quote_signal",
]:
    add_bullet(doc, feature)


# ============================================================
# 4. Exploratory Data Analysis
# ============================================================

add_heading(doc, "4. Exploratory Data Analysis")

doc.add_paragraph(
    "The target distribution is strongly right-skewed. The median posted rate "
    "is approximately $2,031 while the mean is approximately $2,374, indicating "
    "that a relatively small number of high-rate observations increase the mean."
)

add_table(
    doc,
    ["Target Statistic", "Value"],
    [
        ["Minimum", "$57.22"],
        ["Median", "$2,030.76"],
        ["Mean", "$2,373.98"],
        ["Maximum", "$25,533.00"],
    ],
)

doc.add_paragraph(
    "Distance was the strongest individual numerical predictor of posted rate, "
    "with a correlation of approximately 0.91. This is consistent with the "
    "expected relationship between shipment distance and total freight cost."
)

add_heading(doc, "Equipment Distribution", level=2)

add_table(
    doc,
    ["Equipment", "Training Loads"],
    [
        ["Dry Van", "27,202"],
        ["Reefer", "12,045"],
        ["Flatbed", "8,753"],
    ],
)


# ============================================================
# 5. Data Quality
# ============================================================

add_heading(doc, "5. Data Quality and Preprocessing")

doc.add_paragraph(
    "The primary data-quality issues identified were missing numerical values "
    "and a small number of negative weight observations."
)

add_table(
    doc,
    ["Feature", "Missing Values"],
    [
        ["weight", "300"],
        ["market_index", "374"],
    ],
)

doc.add_paragraph(
    "Missing numerical values were handled using median imputation. Categorical "
    "variables were handled using most-frequent imputation followed by one-hot "
    "encoding."
)

doc.add_paragraph(
    "A total of 292 negative-weight observations were identified. Removing these "
    "observations was tested as a sensitivity experiment, but it did not improve "
    "the validation results. They were therefore retained rather than imposing "
    "an unsupported assumption about the underlying data-generation process."
)


# ============================================================
# 6. Validation Strategy
# ============================================================

add_heading(doc, "6. Validation Strategy")

doc.add_paragraph(
    "A chronological validation strategy was used because the final validation "
    "period occurs after the labeled development period. A random split could "
    "allow information from later observations to influence the model-selection "
    "process and would not accurately represent the forecasting setting."
)

add_heading(doc, "Primary Holdout", level=2)

add_table(
    doc,
    ["Training Period", "Validation Period"],
    [
        ["January–September 2025", "October 2025"],
    ],
)

add_heading(doc, "Rolling Validation", level=2)

add_table(
    doc,
    ["Validation Month", "MAE", "RMSE", "R²"],
    [
        ["July", "$122.01", "$633.14", "0.8230"],
        ["August", "$95.37", "$614.15", "0.8266"],
        ["September", "$99.89", "$614.06", "0.8375"],
        ["October", "$103.75", "$648.45", "0.8200"],
        ["Average", "$105.25", "$627.45", "0.8268"],
    ],
)

doc.add_paragraph(
    "The rolling validation results were reasonably stable across months, "
    "supporting the use of the selected model for the later validation period."
)


# ============================================================
# 7. Feature Engineering
# ============================================================

add_heading(doc, "7. Feature Engineering")

doc.add_paragraph(
    "Multiple feature configurations were evaluated. Experiments included raw "
    "shipment and geographic features, date-derived features, log-transformed "
    "distance, market index, quote signal, and equipment information."
)

doc.add_paragraph(
    "An important finding was that additional date and market-signal features "
    "did not consistently improve out-of-time performance. The final model "
    "therefore emphasizes the feature set that provided the strongest and most "
    "stable validation performance."
)

add_heading(doc, "Final Feature Groups", level=2)

for feature in [
    "distance",
    "weight",
    "pickup latitude / longitude",
    "delivery latitude / longitude",
    "pickup location",
    "delivery location",
    "equipment",
]:
    add_bullet(doc, feature)


# ============================================================
# 8. Model Comparison
# ============================================================

add_heading(doc, "8. Model Comparison")

add_table(
    doc,
    ["Model", "MAE", "RMSE", "R²"],
    [
        ["HistGradientBoosting", "$127.34", "$650.67", "0.8188"],
        ["ExtraTrees", "$158.43", "$705.20", "0.7872"],
        ["RandomForest", "$150.09", "$674.67", "0.8052"],
        ["GradientBoosting", "$106.38", "$649.09", "0.8197"],
    ],
)

doc.add_paragraph(
    "Gradient Boosting provided the strongest overall validation performance "
    "and was selected for hyperparameter tuning."
)


# ============================================================
# 9. Hyperparameter Tuning
# ============================================================

add_heading(doc, "9. Hyperparameter Tuning")

doc.add_paragraph(
    "Gradient Boosting was tuned across the number of estimators, learning rate, "
    "tree depth, minimum leaf size, and subsampling fraction."
)

add_table(
    doc,
    ["Configuration", "MAE", "RMSE", "R²"],
    [
        ["300 estimators, depth 3", "$106.38", "$649.09", "0.8197"],
        ["500 estimators, depth 3, leaf 3", "$104.95", "$648.67", "0.8199"],
        ["500 estimators, depth 4, leaf 3", "$103.90", "$648.71", "0.8199"],
        ["700 estimators, depth 4, subsample 0.8", "$103.75", "$648.45", "0.8200"],
    ],
)

add_heading(doc, "Selected Configuration", level=2)

add_code(
    doc,
    """GradientBoostingRegressor(
    n_estimators=700,
    learning_rate=0.03,
    max_depth=4,
    min_samples_leaf=3,
    subsample=0.8,
    loss="huber",
    random_state=42
)"""
)


# ============================================================
# 10. Loss Function Analysis
# ============================================================

add_heading(doc, "10. Loss Function Analysis")

add_table(
    doc,
    ["Loss Function", "MAE", "RMSE", "R²"],
    [
        ["Huber", "$103.75", "$648.45", "0.8200"],
        ["Squared Error", "$122.65", "$654.57", "0.8166"],
        ["Absolute Error", "$104.56", "$648.45", "0.8200"],
    ],
)

doc.add_paragraph(
    "Huber loss was selected because it produced the best MAE while maintaining "
    "strong RMSE and R² performance. Its robustness is particularly useful for "
    "a dataset containing a small number of extreme target values."
)


# ============================================================
# 11. Error and High-Rate Analysis
# ============================================================

add_heading(doc, "11. Error and High-Rate Analysis")

doc.add_paragraph(
    "The majority of predictions were relatively accurate, but a small number "
    "of extreme observations generated disproportionately large errors."
)

add_table(
    doc,
    ["Absolute Error Threshold", "Number of October Loads"],
    [
        ["> $100", "955"],
        ["> $250", "195"],
        ["> $500", "96"],
        ["> $1,000", "67"],
        ["> $2,000", "47"],
        ["> $5,000", "21"],
    ],
)

doc.add_paragraph(
    "Loads above $7,500 represented only approximately 0.32% of the development "
    "data, while loads above $10,000 represented approximately 0.21%."
)

doc.add_paragraph(
    "Historical route behavior was investigated to determine whether these "
    "extreme events could be predicted from previous observations. The analysis "
    "showed that many October high-rate loads occurred on routes whose previous "
    "loads had normal rates. This indicates that the rare extreme regime is "
    "only weakly identifiable from the supplied features."
)

doc.add_paragraph(
    "A separate high-rate classifier achieved a ROC-AUC of approximately 0.72, "
    "but predicted probabilities remained too low to support reliable hard "
    "classification. The final pipeline therefore avoids an additional hard "
    "classification stage."
)


# ============================================================
# 12. Final Model
# ============================================================

add_heading(doc, "12. Final Model and Training")

doc.add_paragraph(
    "After model selection and validation, the selected Gradient Boosting "
    "configuration was retrained using all 48,000 labeled development rows."
)

doc.add_paragraph(
    "The resulting model was applied to every row in the 12,000-row validation "
    "dataset."
)

add_heading(doc, "Final Prediction Distribution", level=2)

add_table(
    doc,
    ["Statistic", "Prediction"],
    [
        ["Number of predictions", "12,000"],
        ["Minimum", "$190.30"],
        ["Median", "$2,034.24"],
        ["Mean", "$2,356.03"],
        ["Maximum", "$6,820.59"],
    ],
)


# ============================================================
# 13. December Scenario
# ============================================================

add_heading(doc, "13. December Prediction Scenario")

doc.add_paragraph(
    "The assessment requires predictions for a fixed shipment scenario across "
    "all 31 days of December 2025."
)

add_table(
    doc,
    ["Input", "Value"],
    [
        ["Pickup", "Lexington"],
        ["Delivery", "Fort Wayne"],
        ["Distance", "360"],
        ["Equipment", "Dry Van"],
        ["Weight", "32,000"],
    ],
)

doc.add_paragraph(
    "The final model produced a prediction of approximately $812.30 for the "
    "fixed scenario across the month."
)

doc.add_paragraph(
    "The flat prediction is expected because the final model does not rely on "
    "date-derived features and the other shipment inputs remain unchanged for "
    "each December row."
)

if CHART.exists():
    add_heading(doc, "Required December Chart", level=2)
    doc.add_picture(str(CHART), width=Inches(6.2))
    last_paragraph = doc.paragraphs[-1]
    last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
else:
    doc.add_paragraph(
        "December chart could not be embedded because the scorer output image "
        "was not found at the expected path."
    )


# ============================================================
# 14. Official Scoring
# ============================================================

add_heading(doc, "14. Official Output Validation")

doc.add_paragraph(
    "The provided scoring script was executed against the final prediction "
    "files."
)

add_code(
    doc,
    "python score.py --predictions outputs/validation_predictions.csv "
    "--december-predictions outputs/december_predictions.csv "
    "--output-dir scorer_results"
)

add_table(
    doc,
    ["Check", "Result"],
    [
        ["Validation prediction rows", "12,000 — Passed"],
        ["Validation prediction columns", "Passed"],
        ["Validation load IDs", "Passed"],
        ["December prediction rows", "31 — Passed"],
        ["December date coverage", "Passed"],
        ["December chart", "Generated"],
    ],
)

doc.add_paragraph(
    "The scoring script does not calculate the hidden evaluation metrics. "
    "Those metrics are calculated by Spotter after submission."
)


# ============================================================
# 15. Reproducibility
# ============================================================

add_heading(doc, "15. Reproducibility")

doc.add_paragraph(
    "The complete solution can be reproduced from the repository using the "
    "following steps."
)

add_code(doc, "python -m venv .venv")
add_code(doc, ".venv\\Scripts\\Activate.ps1")
add_code(doc, "pip install -r requirements.txt")
add_code(doc, "python src\\final_model.py")
add_code(doc, "python src\\december_predictions.py")
add_code(
    doc,
    "python score.py --predictions outputs\\validation_predictions.csv "
    "--december-predictions outputs\\december_predictions.csv "
    "--output-dir scorer_results"
)


# ============================================================
# 16. Limitations
# ============================================================

add_heading(doc, "16. Limitations")

for limitation in [
    "The supplied feature set does not fully explain the rare extreme-rate regime.",
    "A small number of extreme observations disproportionately affect RMSE.",
    "Historical route behavior was not sufficient to reliably predict all high-rate shocks.",
    "The final model intentionally prioritizes stable performance across the dominant pricing regime.",
    "Hidden evaluation performance may differ from the development validation results."
]:
    add_bullet(doc, limitation)


# ============================================================
# 17. Conclusion
# ============================================================

add_heading(doc, "17. Conclusion")

doc.add_paragraph(
    "The final solution uses a time-aware machine learning workflow designed "
    "to reflect the chronological nature of the assessment. Exploratory analysis "
    "identified distance as the strongest individual predictor and revealed a "
    "small but important extreme-rate regime."
)

doc.add_paragraph(
    "After comparing multiple ensemble models and tuning Gradient Boosting, "
    "the selected Huber-loss GradientBoostingRegressor achieved an October "
    "validation MAE of $103.75 and R² of 0.8200, with an average rolling "
    "validation MAE of $105.25."
)

doc.add_paragraph(
    "The final model was retrained on all labeled development data and produced "
    "valid predictions for all 12,000 validation loads. The required 31-day "
    "December scenario was also generated and successfully validated by the "
    "provided scoring script."
)


# ============================================================
# Final document formatting
# ============================================================

for paragraph in doc.paragraphs:
    paragraph.paragraph_format.line_spacing = 1.08

# Add footer
for section in doc.sections:
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = p.add_run("Freight Rate Prediction — Machine Learning Assessment")
    run.font.size = Pt(8)
    run.italic = True


# Save
doc.save(OUTPUT)

print()
print("=" * 70)
print("REPORT GENERATED SUCCESSFULLY")
print("=" * 70)
print(f"Output: {OUTPUT}")
print(f"Chart embedded: {CHART.exists()}")
print("=" * 70)